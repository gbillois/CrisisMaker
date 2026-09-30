"""Builds tests/fixtures/exercise-deck.pptx: a crisis exercise deck shaped like the real ones
(context, objectives, players, phases, a chronogram table over two slides, one inject per
slide, speaker notes, a hidden slide, a chart and a SmartArt), for the PPTX reader tests.

    pip install python-pptx && python3 tests/fixtures/make-exercise-deck.py
"""
import os, re, shutil, tempfile, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'exercise-deck.pptx')

prs = Presentation()
prs.core_properties.title = 'Exercice de crise cyber - Opération Chaîne du froid'
prs.core_properties.author = 'Crisis team'
TITLE, TITLE_BODY, TITLE_ONLY, BLANK = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5], prs.slide_layouts[6]

def bullets(slide, items):
    body = slide.placeholders[1].text_frame
    body.text = ''
    first = True
    for level, text in items:
        p = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        p.text = text
        p.level = level

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def table(slide, rows, top=1.5, merge_col=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.3), Inches(top), Inches(9.4), Inches(0.4 * len(rows)))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            shape.table.cell(r, c).text = value
    if merge_col is not None:
        # The phase column: one cell merged over the rows of the same phase.
        start = 1
        for r in range(2, len(rows) + 1):
            if r == len(rows) or rows[r][merge_col]:
                if r - 1 > start:
                    shape.table.cell(start, merge_col).merge(shape.table.cell(r - 1, merge_col))
                start = r
    return shape

# 1. Cover
s = prs.slides.add_slide(TITLE)
s.shapes.title.text = 'Opération Chaîne du froid'
s.placeholders[1].text = 'Exercice de crise cyber - Groupe Frigolog - 12 mars 2026'

# 2. Agenda (to be recognised as logistics, not as content)
s = prs.slides.add_slide(TITLE_BODY)
s.shapes.title.text = 'Déroulé de la journée'
bullets(s, [(0, '08:30 Accueil des participants'), (0, '09:00 Début de l’exercice (durée 3 h)'), (0, '12:00 RETEX à chaud')])

# 3. Context
s = prs.slides.add_slide(TITLE_BODY)
s.shapes.title.text = 'Contexte'
bullets(s, [(0, 'Frigolog, logisticien du froid, 2 400 salariés, 14 entrepôts en France'),
            (0, 'Activité critique : pilotage des chambres froides via le WMS & la GTB'),
            (1, 'Pic d’activité avant Pâques, clients grande distribution')])
notes(s, 'Rappeler aux joueurs que la société est fictive.')

# 4. Objectives
s = prs.slides.add_slide(TITLE_BODY)
s.shapes.title.text = 'Objectifs de l’exercice'
bullets(s, [(0, 'Tester la mobilisation de la cellule de crise décisionnelle'),
            (0, 'Direction : arbitrer l’arrêt du WMS dans l’incertitude'),
            (0, 'Communication : tenir des messages cohérents face aux médias'),
            (1, 'Juridique : notifier la CNIL sous 72 h')])

# 5. Players: a table of cells
s = prs.slides.add_slide(TITLE_ONLY)
s.shapes.title.text = 'Participants et cellules'
table(s, [['Cellule', 'Participants', 'Rôle'],
          ['Cellule décisionnelle', 'DG, DAF, DRH, Dir. juridique', 'Décide et arbitre'],
          ['Cellule opérationnelle IT', 'DSI, RSSI, SOC', 'Qualifie et remédie'],
          ['Communication', 'Dir. communication', 'Messages internes et externes']])

# 6. Phases: a chevron row inside a group (reading order left to right)
s = prs.slides.add_slide(TITLE_ONLY)
s.shapes.title.text = 'Les phases du scénario'
group = s.shapes.add_group_shape()
phases = [('Phase 1 - Détection', 'H+0:00 → H+0:45'), ('Phase 2 - Escalade', 'H+0:45 → H+2:00'), ('Phase 3 - Sortie de crise', 'H+2:00 → H+3:00')]
# Drawn right to left: the reader must follow the positions, not the drawing order.
for i in reversed(range(len(phases))):
    name, when = phases[i]
    box = group.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(0.3 + 3.1 * i), Inches(2), Inches(3), Inches(1.2))
    box.text_frame.text = name
    box.text_frame.add_paragraph().text = when

# 7. Chronogram, part 1 (merged phase cells)
header = ['Horaire', 'Phase', 'Émetteur', 'Destinataire', 'Canal', 'Contenu']
s = prs.slides.add_slide(TITLE_ONLY)
s.shapes.title.text = 'Chronogramme (1/2)'
table(s, [header,
          ['H+0:00', 'Détection', 'SOC', 'Cellule opérationnelle IT', 'Mail', 'Alertes EDR sur 3 serveurs du WMS'],
          ['H+0:15', '', 'Chef d’entrepôt Rungis', 'Cellule opérationnelle IT', 'Téléphone', 'Les écrans du WMS affichent une note de rançon'],
          ['H+0:40', '', 'DSI', 'Cellule décisionnelle', 'Mail', 'Demande de convocation de la cellule de crise']], merge_col=1)

# 8. Chronogram, part 2 (header repeated)
s = prs.slides.add_slide(TITLE_ONLY)
s.shapes.title.text = 'Chronogramme (2/2)'
table(s, [header,
          ['H+1:00', 'Escalade', 'Journaliste Le Parisien', 'Communication', 'Téléphone', 'Rumeur de denrées avariées, demande de réaction'],
          ['H+1:30', '', 'Client Carrefour', 'Cellule décisionnelle', 'Mail', 'Menace de pénalités si les livraisons s’arrêtent'],
          ['H+2:15', 'Sortie de crise', 'ANSSI', 'Cellule opérationnelle IT', 'Mail', 'Point de situation et recommandations']], merge_col=1)

# 9. One inject per slide
s = prs.slides.add_slide(TITLE_BODY)
s.shapes.title.text = 'Stimulus 7 - Tweet viral'
bullets(s, [(0, 'H+1:10'), (0, 'Émetteur : @FoodWatchFR'), (0, 'Destinataire : Communication'), (0, 'Canal : Réseaux sociaux'),
            (0, '« Des yaourts tièdes livrés chez @Carrefour ? #Frigolog cache quelque chose »')])
notes(s, 'Relance si pas de réaction sous 10 min : un deuxième tweet avec photo.')

# 10. Hidden slide (animator only)
s = prs.slides.add_slide(TITLE_BODY)
s.shapes.title.text = 'Réservé aux animateurs'
bullets(s, [(0, 'Si la cellule décide d’arrêter le WMS, injecter le stimulus 12 bis')])
s._element.set('show', '0')

# 11. Chart: injects per cell
s = prs.slides.add_slide(TITLE_ONLY)
s.shapes.title.text = 'Charge par cellule'
data = CategoryChartData()
data.categories = ['Décisionnelle', 'IT', 'Communication']
data.add_series('Stimuli', (5, 9, 4))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4.5), data)

# 12. SmartArt placeholder slide (the diagram parts are added below)
s = prs.slides.add_slide(TITLE_ONLY)
s.shapes.title.text = 'Chronologie de l’attaque'

# 13. Debrief
s = prs.slides.add_slide(TITLE_BODY)
s.shapes.title.text = 'RETEX et évaluation'
bullets(s, [(0, 'Tour de table des joueurs'), (0, 'Grille d’évaluation par objectif')])

prs.save(OUT)

# SmartArt: python-pptx cannot write one, so the diagram data part is added by hand.
DGM = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>'
       '<dgm:pt modelId="{0}" type="doc"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:p><a:endParaRPr/></a:p></dgm:t></dgm:pt>'
       + ''.join('<dgm:pt modelId="{%d}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:p><a:r><a:t>%s</a:t></a:r></a:p></dgm:t></dgm:pt>' % (i + 1, t)
                 for i, t in enumerate(['J-21 : hameçonnage d’un comptable', 'J-3 : exfiltration de 80 Go', 'Jour J 06:40 : chiffrement du WMS']))
       + '<dgm:pt modelId="{9}" type="sibTrans"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:p><a:r><a:t>ignored</a:t></a:r></a:p></dgm:t></dgm:pt>'
       '</dgm:ptLst><dgm:cxnLst/></dgm:dataModel>')
FRAME = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="90" name="Diagram 1"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
         '<p:xfrm><a:off x="457200" y="1600200"/><a:ext cx="8229600" cy="4525963"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
         '<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:dm="rIdDgm1" r:lo="rIdDgm1" r:qs="rIdDgm1" r:cs="rIdDgm1"/>'
         '</a:graphicData></a:graphic></p:graphicFrame>')
tmp = tempfile.mktemp(suffix='.pptx')
with zipfile.ZipFile(OUT) as src, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():
        data = src.read(item.filename)
        if item.filename == 'ppt/slides/slide12.xml':
            data = data.decode('utf8').replace('</p:spTree>', FRAME + '</p:spTree>').encode('utf8')
        elif item.filename == 'ppt/slides/_rels/slide12.xml.rels':
            data = data.decode('utf8').replace('</Relationships>', '<Relationship Id="rIdDgm1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData" Target="../diagrams/data1.xml"/></Relationships>').encode('utf8')
        elif item.filename == '[Content_Types].xml':
            data = data.decode('utf8').replace('</Types>', '<Override PartName="/ppt/diagrams/data1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml"/></Types>').encode('utf8')
        dst.writestr(item, data)
    dst.writestr('ppt/diagrams/data1.xml', DGM)
shutil.move(tmp, OUT)
print('wrote', OUT)
