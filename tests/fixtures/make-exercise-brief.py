"""Builds tests/fixtures/exercise-brief.docx: an exercise proposal written in Word (headings,
bullets, a table), for the document reader tests.

    pip install python-docx && python3 tests/fixtures/make-exercise-brief.py
"""
import os
from docx import Document

doc = Document()
doc.core_properties.title = 'Proposal: cyber crisis exercise for Northwind Health'
doc.add_heading('Cyber crisis exercise proposal', 0)
doc.add_paragraph('Prepared for Northwind Health, regional hospital group.')
doc.add_heading('Context', 1)
doc.add_paragraph('Northwind Health runs 6 hospitals and 40 clinics. Its EHR was migrated to a hosted platform in 2025.')
doc.add_heading('Exercise objectives', 1)
for text in ['Executives: decide on isolating the EHR under uncertainty', 'Communication: brief patients and the press consistently', 'Legal: notify the regulator within 72 hours']:
    doc.add_paragraph(text, style='List Bullet')
doc.add_heading('Participants', 1)
table = doc.add_table(rows=3, cols=2)
for r, row in enumerate([['Cell', 'Players'], ['Executive crisis cell', 'CEO, CMO, CFO, General counsel'], ['IT crisis cell', 'CIO, CISO, SOC lead']]):
    for c, value in enumerate(row):
        table.cell(r, c).text = value
doc.add_heading('Approach and budget', 1)
doc.add_paragraph('Two design workshops, one 3-hour exercise, a debrief report. 12 days in total.')
doc.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'exercise-brief.docx'))
